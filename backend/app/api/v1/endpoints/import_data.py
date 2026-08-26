import csv
import io
import json
from datetime import date
from typing import Optional

from fastapi import APIRouter, Depends, File, Form, HTTPException, Response, UploadFile, status
from openpyxl import Workbook
from pydantic import BaseModel
from sqlmodel import Session

from ....auth.permissions import get_current_user, user_has_right
from ....db.database import getSession
from ....models.user import User
from ....services.import_service import (
    TABLES,
    build_export,
    build_preview,
    get_schema,
    parse_upload,
    run_commit,
)

router = APIRouter()


@router.get("/schema")
def import_schema(_user: User = Depends(get_current_user)):
    return get_schema()


class ImportRowIn(BaseModel):
    row_number: Optional[int] = None
    values: dict[str, str] = {}


class ImportTableIn(BaseModel):
    entity: str
    sheet_name: Optional[str] = None
    mapping: dict[str, str] = {}
    rows: list[ImportRowIn] = []


class ImportCommitIn(BaseModel):
    duplicate_mode: str = "skip"
    tables: list[ImportTableIn] = []


@router.post("/preview")
async def preview_import(
    file: UploadFile = File(...),
    hints: Optional[str] = Form(None),
    _user: User = Depends(get_current_user),
    session: Session = Depends(getSession),
):
    data = await file.read()
    if not data:
        raise HTTPException(status_code=400, detail="Leë lêer")
    try:
        sheets = parse_upload(file.filename or "", data)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc))
    hint_map: dict[str, str] = {}
    if hints:
        try:
            parsed = json.loads(hints)
            if isinstance(parsed, dict):
                hint_map = {str(k): str(v) for k, v in parsed.items()}
        except json.JSONDecodeError:
            pass
    return build_preview(session, sheets, hint_map)


@router.post("/commit")
def commit_import(
    body: ImportCommitIn,
    user: User = Depends(get_current_user),
    session: Session = Depends(getSession),
):
    if not body.tables:
        raise HTTPException(status_code=400, detail="Geen taballe om in te voer nie")

    missing = []
    for t in body.tables:
        spec = TABLES.get(t.entity)
        if spec is None:
            raise HTTPException(status_code=400, detail=f"Onbekende tabel: {t.entity}")
        if not user_has_right(session, user.role_id, spec.right):
            missing.append(spec.right)
    if missing:
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail=f"Ontbrekende regte: {', '.join(sorted(set(missing)))}",
        )

    if body.duplicate_mode not in ("skip", "update", "create"):
        raise HTTPException(status_code=400, detail="duplicate_mode moet skip, update of create wees")

    try:
        result = run_commit(
            session,
            {"duplicate_mode": body.duplicate_mode, "tables": [t.model_dump() for t in body.tables]},
            actor_id=user.user_id,
        )
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc))
    return result


class ExportTableIn(BaseModel):
    entity: str
    columns: Optional[list[str]] = None


class ExportIn(BaseModel):
    format: str = "xlsx"
    template: bool = False
    tables: list[ExportTableIn] = []


@router.post("/export")
def export_records(
    body: ExportIn,
    user: User = Depends(get_current_user),
    session: Session = Depends(getSession),
):
    if body.format not in ("csv", "xlsx"):
        raise HTTPException(status_code=400, detail="format moet csv of xlsx wees")
    if not body.tables:
        raise HTTPException(status_code=400, detail="Geen tabelle gekies nie")
    if body.format == "csv" and len(body.tables) > 1:
        raise HTTPException(status_code=400, detail="CSV kan slegs een tabel bevat — kies XLSX vir meer as een")

    missing_rights = []
    for t in body.tables:
        spec = TABLES.get(t.entity)
        if spec is None:
            raise HTTPException(status_code=400, detail=f"Onbekende tabel: {t.entity}")
        if not user_has_right(session, user.role_id, spec.right):
            missing_rights.append(spec.right)
        if t.columns is not None:
            valid = set(spec.all_targets())
            if not t.columns:
                raise HTTPException(status_code=400, detail=f"Geen kolomme gekies vir {spec.label} nie")
            unknown = [c for c in t.columns if c not in valid]
            if unknown:
                raise HTTPException(
                    status_code=400,
                    detail=f"Onbekende kolomme vir {spec.label}: {', '.join(unknown)}",
                )
    if missing_rights:
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail=f"Ontbrekende regte: {', '.join(sorted(set(missing_rights)))}",
        )

    tables = build_export(session, [t.model_dump() for t in body.tables], template=body.template)
    stamp = date.today().strftime("%Y%m%d")
    names = "-".join(t.entity for t in body.tables)

    if body.format == "csv":
        table = tables[0]
        buf = io.StringIO()
        writer = csv.writer(buf)
        writer.writerow(table["columns"])
        for row in table["rows"]:
            writer.writerow(row)
        filename = f"uitvoer-{names}-{stamp}.csv"
        return Response(
            content=buf.getvalue().encode("utf-8-sig"),
            media_type="text/csv; charset=utf-8",
            headers={"Content-Disposition": f'attachment; filename="{filename}"'},
        )

    wb = Workbook()
    wb.remove(wb.active)
    used_titles = set()
    for table in tables:
        title = table["sheet_name"]
        base, i = title, 2
        while title in used_titles:
            title = f"{base}_{i}"
            i += 1
        used_titles.add(title)
        ws = wb.create_sheet(title=title)
        ws.append(table["columns"])
        for row in table["rows"]:
            ws.append(row)
    out = io.BytesIO()
    wb.save(out)
    filename = f"uitvoer-{names}-{stamp}.xlsx"
    return Response(
        content=out.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
