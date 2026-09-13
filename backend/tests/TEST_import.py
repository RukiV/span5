"""Tests vir die CSV/XLSX data-invoer (import) funksionaliteit.

Dek: lêer-parsing, blad→tabel herkenning, kolom-outo-mapping, invoer-volgorde
(ouers voor kinders), verwysings-resolusie met fout-rye, duplikaat hantering
(skip/update/create), gebruiker-import (rol afgedwing na "User"), en regte.
"""

import csv
import io

import pytest
from openpyxl import Workbook


def _make_xlsx(sheets: dict) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(title=name)
        for row in rows:
            ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _preview(client, headers, filename, content, hints=None):
    data = {"file": (filename, content)}
    if hints:
        data["hints"] = hints
    return client.post("/api/v1/import/preview", files=data, headers=headers)


def _commit(client, headers, duplicate_mode, tables):
    return client.post(
        "/api/v1/import/commit",
        json={"duplicate_mode": duplicate_mode, "tables": tables},
        headers=headers,
    )


def _table_from_preview(preview_json, sheet_name):
    sheet = next(s for s in preview_json["sheets"] if s["sheet_name"] == sheet_name)
    rows = [
        {"row_number": r["row_number"], "values": vals}
        for r, vals in zip(sheet["rows"], _last_values[0])
    ]
    return {
        "entity": sheet["entity"],
        "sheet_name": sheet_name,
        "mapping": sheet["mapping"],
        "rows": rows,
    }


_last_values = [{}]


def _capture_values(preview_json, sheet_name):
    sheet = next(s for s in preview_json["sheets"] if s["sheet_name"] == sheet_name)
    _last_values[0] = [r["values"] for r in sheet["rows"]]
    return sheet


ADMIN = None


@pytest.fixture(name="admin_headers")
def admin_headers_fixture(headers_for):
    return headers_for("admin")


def test_full_hierarchy_xlsx_import_order_and_refs(client, admin_headers):
    xlsx = _make_xlsx({
        "Kamers": [["Naam", "Kode", "Gebou"], ["A101", "A101", "Hoofgebou"]],
        "Geboue": [["Gebou Naam", "Tipe", "Kampus"], ["Hoofgebou", "Onderwys", "Main"]],
        "Kampe": [["Naam", "Tipe", "Straatnommer", "Straatnaam"], ["Main", "Kampus", "1", "Kerkstraat"]],
    })
    resp = _preview(client, admin_headers, "fasiliteite.xlsx", xlsx)
    assert resp.status_code == 200, resp.text
    preview = resp.json()

    by_sheet = {s["sheet_name"]: s for s in preview["sheets"]}
    assert by_sheet["Kampe"]["entity"] == "location"
    assert by_sheet["Geboue"]["entity"] == "building"
    assert by_sheet["Kamers"]["entity"] == "room"
    assert by_sheet["Kampe"]["missing_required"] == []
    assert by_sheet["Kampe"]["counts"]["error"] == 0
    assert by_sheet["Kamers"]["rows"][0]["status"] == "create"

    tables = []
    for name in ("Kampe", "Geboue", "Kamers"):
        sheet = by_sheet[name]
        tables.append({
            "entity": sheet["entity"],
            "sheet_name": name,
            "mapping": sheet["mapping"],
            "rows": [{"row_number": r["row_number"], "values": r["values"]} for r in sheet["rows"]],
        })

    resp = _commit(client, admin_headers, "skip", tables)
    assert resp.status_code == 200, resp.text
    result = resp.json()
    assert result["totals"]["create"] == 3
    assert result["totals"]["error"] == 0

    locs = client.get("/api/v1/location", headers=admin_headers).json()
    assert any(l["location_name"] == "Main" for l in locs)
    blds = client.get("/api/v1/building", headers=admin_headers).json()
    main_bld = next(b for b in blds if b["building_name"] == "Hoofgebou")
    main_loc = next(l for l in locs if l["location_name"] == "Main")
    assert main_bld["location_id"] == main_loc["location_id"]
    rooms = client.get("/api/v1/rooms", headers=admin_headers).json()
    room = next(r for r in rooms if r.get("room_code") == "A101" or r["room_name"] == "A101")
    assert room["building_id"] == main_bld["building_id"]


def test_duplicate_modes_skip_update_create(client, admin_headers):
    xlsx = _make_xlsx({
        "Kampe": [
            ["Naam", "Tipe", "Straatnommer", "Straatnaam"],
            ["AanCampus", "Kampus", "10", "Straatweg"],
        ],
    })
    resp = _preview(client, admin_headers, "k.xlsx", xlsx)
    sheet = _capture_values(resp.json(), "Kampe")
    table = [{
        "entity": sheet["entity"], "sheet_name": "Kampe",
        "mapping": sheet["mapping"],
        "rows": [{"row_number": r["row_number"], "values": r["values"]}
                 for r in sheet["rows"]],
    }]

    first = _commit(client, admin_headers, "skip", table).json()
    assert first["totals"]["create"] == 1

    second_skip = _commit(client, admin_headers, "skip", table).json()
    assert second_skip["totals"]["skip"] == 1
    assert second_skip["totals"]["create"] == 0

    update_tbl = dict(table[0])
    update_tbl["rows"] = [dict(table[0]["rows"][0], values={
        **table[0]["rows"][0]["values"], "Straatnaam": "Nuwe Straat"})]
    second_update = _commit(client, admin_headers, "update", [update_tbl]).json()
    assert second_update["totals"]["update"] == 1

    third_create = _commit(client, admin_headers, "create", table).json()
    assert third_create["totals"]["create"] == 1

    locs = client.get("/api/v1/location", headers=admin_headers).json()
    matches = [l for l in locs if l["location_name"] == "AanCampus"]
    assert len(matches) == 2
    streets = {l["location_streetname"] for l in matches}
    assert streets == {"Nuwe Straat", "Straatweg"}


def test_missing_reference_fails_row_only(client, admin_headers):
    xlsx = _make_xlsx({
        "Aktiwiteite": [
            ["Naam", "Handelsmerk", "Serienommer", "Tipe", "Kamer"],
            ["Projektor X", "Epson", "SN001", "Rekenaar", "A101"],
            ["Beker", "Russell", "SN002", "OnbekendTipe", ""],
        ],
    })
    resp = _preview(client, admin_headers, "a.xlsx", xlsx)
    sheet = _capture_values(resp.json(), "Aktiwiteite")
    assert resp.status_code == 200
    rows_by_num = {r["row_number"]: r for r in sheet["rows"]}
    assert rows_by_num[sheet["rows"][0]["row_number"]]["status"] == "error"

    table = [{
        "entity": "asset", "sheet_name": "Aktiwiteite",
        "mapping": sheet["mapping"],
        "rows": [{"row_number": r["row_number"], "values": r["values"]}
                 for r in sheet["rows"]],
    }]
    result = _commit(client, admin_headers, "skip", table).json()
    assert result["totals"]["create"] == 0
    assert result["totals"]["error"] == 2
    assets_now = client.get("/api/v1/assets", headers=admin_headers).json()
    assert len([a for a in assets_now if a["asset_serial"] in ("SN001", "SN002")]) == 0


def test_asset_resolves_room_and_type_created_same_run(client, admin_headers):
    xlsx = _make_xlsx({
        "Kampe": [["Naam", "Tipe", "Straatnommer", "Straatnaam"], ["C Campus", "Kampus", "3", "Weg"]],
        "Geboue": [["Naam", "Kampus"], ["C Gebou", "C Campus"]],
        "Kamers": [["Naam", "Kode", "Gebou"], ["Lab 1", "L001", "C Gebou"]],
        "AssetTypes": [["Naam"], ["Rekenaar"]],
        "Aktiwiteite": [
            ["Naam", "Handelsmerk", "Serienommer", "Tipe", "Kamer Kode"],
            ["HP EliteDesk", "HP", "SN777", "Rekenaar", "L001"],
        ],
    })
    resp = _preview(client, admin_headers, "alles.xlsx", xlsx)
    preview = resp.json()
    by_sheet = {s["sheet_name"]: s for s in preview["sheets"]}
    tables = [
        {"entity": s["entity"], "sheet_name": n, "mapping": s["mapping"],
         "rows": [{"row_number": r["row_number"], "values": r["values"]} for r in s["rows"]]}
        for n, s in by_sheet.items() if s["entity"]
    ]
    result = _commit(client, admin_headers, "skip", tables).json()
    assert result["totals"]["error"] == 0, result
    assert result["totals"]["create"] == 5

    assets = client.get("/api/v1/assets", headers=admin_headers).json()
    hp = next(a for a in assets if a["asset_serial"] == "SN777")
    rooms = client.get("/api/v1/rooms", headers=admin_headers).json()
    lab = next(r for r in rooms if r["room_code"] == "L001")
    assert hp["room_id"] == lab["room_id"]
    types = client.get("/api/v1/assettypes", headers=admin_headers).json()
    atype = next(t for t in types if t["assettype_name"] == "Rekenaar")
    assert hp["assettype_id"] == atype["assettype_id"]


def test_user_import_forced_role_and_unique_email(client, admin_headers):
    xlsx = _make_xlsx({
        "Gebruikers": [
            ["Voornaam", "Van", "E-pos", "Status"],
            ["Pieter", "Botha", "pieter@import.local", "Aktief"],
        ],
    })
    resp = _preview(client, admin_headers, "g.xlsx", xlsx)
    sheet = _capture_values(resp.json(), "Gebruikers")
    table = [{
        "entity": "user", "sheet_name": "Gebruikers",
        "mapping": sheet["mapping"],
        "rows": [{"row_number": r["row_number"], "values": r["values"]} for r in sheet["rows"]],
    }]
    res = _commit(client, admin_headers, "skip", table).json()
    assert res["totals"]["create"] == 1

    users = client.get("/api/v1/users", headers=admin_headers).json()
    piet = next(u for u in users if u["user_email"] == "pieter@import.local")
    assert piet["role_id"] == 1

    again_create = _commit(client, admin_headers, "create", table).json()
    assert again_create["totals"]["error"] == 1

    again_update = _commit(client, admin_headers, "update", table).json()
    assert again_update["totals"]["update"] == 1


def test_fault_duplicate_matching_ignores_closed(client, admin_headers):
    xlsx = _make_xlsx({
        "Foute": [["Beskrywing", "Prioriteit"], ["Projector kapot in A101", "Hoog"]],
    })
    resp = _preview(client, admin_headers, "f.xlsx", xlsx)
    sheet = _capture_values(resp.json(), "Foute")
    table = [{
        "entity": "fault", "sheet_name": "Foute",
        "mapping": sheet["mapping"],
        "rows": [{"row_number": r["row_number"], "values": r["values"]} for r in sheet["rows"]],
    }]
    res = _commit(client, admin_headers, "skip", table).json()
    assert res["totals"]["create"] == 1

    faults = client.get("/api/v1/fault", headers=admin_headers).json()
    fid = next(f for f in faults if "Projector kapot" in f["fault_description"])["fault_id"]

    close = client.patch(f"/api/v1/fault/{fid}", json={"fault_status": "Gesluit"},
                         headers=admin_headers)
    assert close.status_code == 200, close.text

    reimport = _commit(client, admin_headers, "skip", table).json()
    assert reimport["totals"]["create"] == 1, "closed fault moet nie as duplikaat tel nie"


def test_csv_upload_and_commit(client, admin_headers):
    csv_bytes = "Naam;Tipe;Straatnommer;Straatnaam\nCSV Campus;Kampus;5;Testweg\n".encode("utf-8")
    resp = _preview(client, admin_headers, "kampe.csv", csv_bytes)
    assert resp.status_code == 200, resp.text
    sheet = _capture_values(resp.json(), "kampe")
    assert sheet["entity"] == "location"

    result = _commit(client, admin_headers, "skip", [{
        "entity": sheet["entity"], "sheet_name": "kampe",
        "mapping": sheet["mapping"],
        "rows": [{"row_number": r["row_number"], "values": r["values"]} for r in sheet["rows"]],
    }]).json()
    assert result["totals"]["create"] == 1


def test_commit_requires_right(client, headers_for):
    student = headers_for("student")
    body = {"duplicate_mode": "skip", "tables": [{
        "entity": "location", "sheet_name": "x", "mapping": {}, "rows": [],
    }]}
    resp = client.post("/api/v1/import/commit", json=body, headers=student)
    assert resp.status_code == 403


def test_unsupported_file_rejected(client, admin_headers):
    resp = _preview(client, admin_headers, "data.docx", b"nonsense")
    assert resp.status_code == 400


# ------------------------- UITVOER (export) -------------------------

def _seed_hierarchy(client, headers):
    """Saai Kampe→Geboue→Kamers→Tipe→Aktiwiteit deur die invoer self."""
    xlsx = _make_xlsx({
        "Kampe": [["Naam", "Tipe", "Straatnommer", "Straatnaam"], ["R Campus", "Kampus", "9", "Uitweg"]],
        "Geboue": [["Naam", "Tipe", "Kampus"], ["R Gebou", "Onderwys", "R Campus"]],
        "Kamers": [["Naam", "Kode", "Gebou"], ["Lab R", "R001", "R Gebou"]],
        "AssetTypes": [["Naam"], ["Skerm"]],
        "Aktiwiteite": [
            ["Naam", "Handelsmerk", "Serienommer", "Status", "Buitelug", "Aanskafdatum", "Tipe", "Kamer Kode"],
            ["Monitor A", "Dell", "SN-EX-1", "Aktief", "Nee", "2024-03-01 09:30:00", "Skerm", "R001"],
        ],
    })
    resp = _preview(client, headers, "saad.xlsx", xlsx)
    by_sheet = {s["sheet_name"]: s for s in resp.json()["sheets"]}
    tables = [
        {"entity": s["entity"], "sheet_name": n, "mapping": s["mapping"],
         "rows": [{"row_number": r["row_number"], "values": r["values"]} for r in s["rows"]]}
        for n, s in by_sheet.items() if s["entity"]
    ]
    result = _commit(client, headers, "skip", tables).json()
    assert result["totals"]["error"] == 0, result
    assert result["totals"]["create"] == 5


def test_job_quotes_link_via_import_and_export(client, admin_headers):
    """'Kwotasies'-kolom op Werkskaarte koppel kwotasies (datum;e-pos|...) aan die job."""
    _seed_hierarchy(client, admin_headers)

    xlsx = _make_xlsx({
        "Gebruikers": [
            ["Voornaam", "Van", "E-pos", "Status"],
            ["Daan", "Smit", "daan@kampus.co.za", "Aktief"],
        ],
        "Kotasies": [
            ["Datum", "Status", "Kontrakteur (e-pos)"],
            ["2025-06-01", "Nuut", "daan@kampus.co.za"],
            ["2025-06-02", "Nuut", "daan@kampus.co.za"],
        ],
        "Werkskaarte": [
            ["Beskrywing", "Status", "Kwotasies"],
            ["Vervang projektorlamp", "Geskeduleer",
             "2025-06-01;daan@kampus.co.za|2025-06-02;daan@kampus.co.za"],
        ],
    })
    resp = _preview(client, admin_headers, "kwota.xlsx", xlsx)
    by_sheet = {s["sheet_name"]: s for s in resp.json()["sheets"]}
    tables = [
        {"entity": s["entity"], "sheet_name": n, "mapping": s["mapping"],
         "rows": [{"row_number": r["row_number"], "values": r["values"]} for r in s["rows"]]}
        for n, s in by_sheet.items() if s["entity"]
    ]
    result = _commit(client, admin_headers, "skip", tables).json()
    assert result["totals"]["error"] == 0, result

    # Die job het albei kwotasie-id's gekry (eerste = primêre keuse)
    jobs = client.get("/api/v1/job", headers=admin_headers).json()
    job_row = next(j for j in jobs if j["job_desc"] == "Vervang projektorlamp")
    quotes = client.get("/api/v1/quotes", headers=admin_headers).json()
    qids = {q["quote_date"][:10]: q["quote_id"] for q in quotes
            if q.get("contractor_id") is not None}
    assert qids, quotes
    first, second = sorted(qids.values())
    assert job_row["quote_ids"] == f"{first},{second}"
    assert job_row["quote_id"] == first

    # Uitvoer render dieselfde "datum;e-pos|..." formaat terug
    exp = client.post("/api/v1/import/export", json={
        "format": "csv",
        "tables": [{"entity": "job",
                    "columns": ["f:job_desc", "f:quote_ids"]}],
    }, headers=admin_headers)
    assert exp.status_code == 200, exp.text
    lines = exp.content.decode("utf-8-sig").strip().splitlines()
    rows = list(csv.reader(lines))
    data = {d["Beskrywing"]: d["Kwotasies"] for d in
            (dict(zip(rows[0], r)) for r in rows[1:])}
    expected = "2025-06-01;daan@kampus.co.za|2025-06-02;daan@kampus.co.za"
    assert data["Vervang projektorlamp"] == expected


def test_export_roundtrip_reimports_cleanly(client, admin_headers):
    _seed_hierarchy(client, admin_headers)

    resp = client.post("/api/v1/import/export", json={
        "format": "xlsx",
        "tables": [{"entity": e} for e in ("location", "building", "room", "assettype", "asset")],
    }, headers=admin_headers)
    assert resp.status_code == 200, resp.text
    assert "uitvoer-" in resp.headers.get("content-disposition", "")

    # Die uitvoer-lêer word direk teruggevoer:
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(resp.content))
    assert wb.sheetnames == ["Terreine", "Geboue", "Lokale", "Aktiwiteitstipes", "Bates"]
    buf = io.BytesIO()
    wb.save(buf)

    preview = _preview(client, admin_headers, "ronde.xlsx", buf.getvalue()).json()
    expected_entities = {
        "Terreine": ("location", 1), "Geboue": ("building", 1), "Lokale": ("room", 1),
        "Aktiwiteitstipes": ("assettype", 1), "Bates": ("asset", 1),
    }
    for sheet_name, (entity, row_count) in expected_entities.items():
        sheet = next(s for s in preview["sheets"] if s["sheet_name"] == sheet_name)
        assert sheet["entity"] == entity, (sheet_name, sheet)
        assert sheet["unmapped_columns"] == [], (sheet_name, sheet["unmapped_columns"])
        assert sheet["missing_required"] == [], (sheet_name, sheet["missing_required"])
        assert sheet["counts"]["error"] == 0, sheet["rows"]
        assert sheet["counts"]["skip"] == row_count, (sheet_name, sheet["counts"])
        assert sheet["counts"]["create"] == 0


def test_export_csv_values_are_import_format(client, admin_headers):
    _seed_hierarchy(client, admin_headers)
    resp = client.post("/api/v1/import/export", json={
        "format": "csv",
        "tables": [{
            "entity": "asset",
            "columns": ["f:asset_serial", "f:asset_isoutdoor", "f:asset_created_datetime", "f:asset_status", "r:room_ref"],
        }],
    }, headers=admin_headers)
    assert resp.status_code == 200, resp.text
    text = resp.content.decode("utf-8-sig")
    lines = text.strip().splitlines()
    # Kanonieke register-volgorde: Naam, Handelsmerk, Serienommer, Status, Buite, Geskep, ⟶Tipe, ⟶Lokaal
    assert lines[0] == "Serienommer,Status,Buite,Geskep,Lokaal"
    assert len(lines) == 2
    data_row = next(csv.reader(lines[1:]))
    assert data_row[0] == "SN-EX-1"
    assert data_row[1] == "Aktief"
    assert data_row[2] == "Nee"
    assert data_row[3] == "2024-03-01 09:30:00"
    assert data_row[4] == "R001"  # kamer volgens kode


def test_export_template_headers_only(client, admin_headers):
    resp = client.post("/api/v1/import/export", json={
        "format": "xlsx", "template": True,
        "tables": [{"entity": "building"}],
    }, headers=admin_headers)
    assert resp.status_code == 200
    from openpyxl import load_workbook
    ws = load_workbook(io.BytesIO(resp.content)).active
    assert ws.title == "Geboue"
    assert [c.value for c in ws[1]] == ["Naam", "Tipes", "Terrein"]
    assert ws.max_row == 1


def test_export_validation_errors(client, admin_headers):
    base = "/api/v1/import/export"
    # onbekende kolom
    r = client.post(base, json={"format": "xlsx", "tables": [
        {"entity": "location", "columns": ["f:location_name", "f:bestaan_nie"]},
    ]}, headers=admin_headers)
    assert r.status_code == 400
    # leë kolomlys
    r = client.post(base, json={"format": "xlsx", "tables": [
        {"entity": "location", "columns": []},
    ]}, headers=admin_headers)
    assert r.status_code == 400
    # csv met meer as een tabel
    r = client.post(base, json={"format": "csv", "tables": [
        {"entity": "location"}, {"entity": "room"},
    ]}, headers=admin_headers)
    assert r.status_code == 400
    # onbekende entiteit
    r = client.post(base, json={"format": "xlsx", "tables": [{"entity": "ruimteskip"}]},
                    headers=admin_headers)
    assert r.status_code == 400


def test_export_requires_right(client, headers_for):
    student = headers_for("student")
    r = client.post("/api/v1/import/export", json={"format": "xlsx", "tables": [
        {"entity": "location"},
    ]}, headers=student)
    assert r.status_code == 403
